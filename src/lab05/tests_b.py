import csv
from pathlib import Path
from openpyxl import load_workbook
from src.lab05.csv_xlsx import csv_to_xlsx


def main() -> None:
    samples = Path('data/samples/lab05_tests_b')
    results = Path('data/out/lab05_tests_b')
    samples.mkdir(parents=True, exist_ok=True)
    results.mkdir(parents=True, exist_ok=True)
    total = 0

    rows = [['name', 'age', 'city'],
            ['Анна', '19', ''], ['Иван, Иванов', '', 'Москва']]
    with (samples / 'people.csv').open('w', encoding='utf-8', newline='') as file:
        csv.writer(file).writerows(rows)

    csv_to_xlsx(str(samples / 'people.csv'), str(results / 'people.xlsx'))
    book = load_workbook(results / 'people.xlsx')
    assert book.sheetnames == ['Sheet1']
    sheet = book.active
    assert list(sheet.values) == [
        ('name', 'age', 'city'), ('Анна', '19', None),
        ('Иван, Иванов', None, 'Москва')]
    for column in sheet.columns:
        width = sheet.column_dimensions[column[0].column_letter].width
        assert width >= 8
        for cell in column:
            if cell.value is not None:
                assert width >= len(str(cell.value)) + 2
    book.close()
    print('CSV -> XLSX: Sheet1, значения и ширина столбцов — OK')
    total += 1

    csv_cases = [('empty', ''), ('header_only', 'name,age\n'),
                 ('duplicate_header', 'name,name\nАнна,19\n'),
                 ('short_row', 'name,age\nАнна\n'),
                 ('empty_header', ',age\nАнна,19\n')]
    for name, text in csv_cases:
        path = samples / (name + '.csv')
        path.write_text(text, encoding='utf-8')
        try:
            csv_to_xlsx(str(path), str(results / 'error.xlsx'))
        except ValueError:
            print(f'{name}.csv: ValueError — OK')
            total += 1
        else:
            raise AssertionError(f'{name}.csv: ожидался ValueError')

    try:
        csv_to_xlsx(str(samples / 'people.csv'), str(results / 'wrong.txt'))
    except ValueError:
        print('Неверное расширение: ValueError — OK')
        total += 1
    else:
        raise AssertionError('Неверное расширение: ожидался ValueError')

    try:
        csv_to_xlsx(str(samples / 'missing.csv'), str(results / 'missing.xlsx'))
    except FileNotFoundError:
        print('Файл отсутствует: FileNotFoundError — OK')
        total += 1
    else:
        raise AssertionError('Отсутствующий файл: ожидался FileNotFoundError')

    print(f'Задание B: {total} проверок. Все пройдены.')


if __name__ == '__main__':
    main()
